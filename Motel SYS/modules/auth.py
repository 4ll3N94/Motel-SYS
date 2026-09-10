"""
===============================================================================
SISTEMA DE GESTIÓN DE MOTELES - MOTEL PRO v2.0
Módulo: modules/auth.py (GESTIÓN CON BOTÓN DE CONTRASEÑA Y AUDITORÍA AVANZADA)
===============================================================================
"""

import os
import openpyxl
import customtkinter as ctk
from tkinter import messagebox
from tkcalendar import DateEntry
from fpdf import FPDF
from datetime import datetime

from database.db_manager import (
    obtener_conexion, hash_clave, obtener_todos_usuarios, actualizar_usuario,
    registrar_auditoria, obtener_logs_auditoria, obtener_ruta_recurso
)


class PDFReporteAuditoria(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)

    def header(self):
        logo = obtener_ruta_recurso("logo_hotel.png")
        if os.path.exists(logo):
            self.image(logo, x=10, y=8, w=20)

        self.set_font("Arial", "B", 12)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN", 0, 1, "C")
        self.set_font("Arial", "B", 9)
        self.cell(0, 4.5, "RIF: J-30250227-6 | INFORME DE AUDITORÍA Y TRAZABILIDAD DE USUARIOS", 0, 1, "C")
        self.set_y(26)

        self.set_fill_color(220, 230, 242)
        self.set_font("Arial", "B", 8)
        self.cell(32, 7, "FECHA / HORA", 1, 0, "C", fill=True)
        self.cell(24, 7, "USUARIO", 1, 0, "C", fill=True)
        self.cell(20, 7, "ROL", 1, 0, "C", fill=True)
        self.cell(35, 7, "ACCIÓN", 1, 0, "C", fill=True)
        self.cell(166, 7, "DESCRIPCIÓN DETALLADA DE LA OPERACIÓN", 1, 0, "L", fill=True)
        self.ln()

    def footer(self):
        self.set_y(-12)
        self.set_font("Arial", "I", 8)
        self.cell(0, 8, f"Auditoría de Seguridad - Página {self.page_no()}", 0, 0, "C")


class AuthManager:
    @staticmethod
    def validar_login(usuario: str, clave: str) -> dict | None:
        if not usuario or not clave:
            messagebox.showwarning("Atención", "Ingrese usuario y contraseña.")
            return None

        clave_hash = hash_clave(clave)
        conn = obtener_conexion()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT id, cedula, nombre, apellido, usuario, rol FROM usuarios WHERE usuario = ? AND clave = ?;", (usuario.strip(), clave_hash))
            user = cursor.fetchone()
            if user:
                user_dict = dict(user)
                registrar_auditoria(user_dict["usuario"], user_dict["rol"], "LOGIN", "Inicio de sesión exitoso en el sistema.")
                return user_dict
            else:
                registrar_auditoria(usuario.strip(), "DESCONOCIDO", "LOGIN_FALLIDO", f"Intento de acceso fallido con el usuario '{usuario}'.")
                messagebox.showerror("Error de Acceso", "Usuario o contraseña incorrectos.")
                return None
        finally:
            conn.close()


class FrameUsuarios(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, fg_color="#191715")
        self.controller = controller

        self.tab_selector = ctk.CTkSegmentedButton(
            self, values=["👥 Gestión de Usuarios", "🛡️ Auditoría de Personal"],
            selected_color="#D4A343", selected_hover_color="#B8892E",
            command=self.cambiar_pestana
        )
        self.tab_selector.set("👥 Gestión de Usuarios")
        self.tab_selector.pack(padx=15, pady=10)

        self.container_pestanas = ctk.CTkFrame(self, fg_color="transparent")
        self.container_pestanas.pack(fill="both", expand=True, padx=10, pady=5)

        self.usuario_id_edicion = None
        self.ver_clave_activa = False
        self.construir_pestana_usuarios()

    def cambiar_pestana(self, pestana):
        for w in self.container_pestanas.winfo_children():
            w.destroy()

        if pestana == "👥 Gestión de Usuarios":
            self.construir_pestana_usuarios()
        else:
            self.construir_pestana_auditoria()

    # =========================================================================
    # PESTAÑA 1: GESTIÓN CON BOTÓN DE OJITO (VER/OCULTAR CONTRASEÑA)
    # =========================================================================
    def construir_pestana_usuarios(self):
        self.container_pestanas.grid_columnconfigure(1, weight=1)
        self.container_pestanas.grid_rowconfigure(0, weight=1)

        self.frame_form = ctk.CTkFrame(self.container_pestanas, fg_color="#23201C", corner_radius=12, border_width=1, border_color="#36322C")
        self.frame_form.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        self.lbl_titulo_form = ctk.CTkLabel(self.frame_form, text="Gestión de Personal", font=("Georgia", 16, "bold"), text_color="#D4A343")
        self.lbl_titulo_form.pack(pady=10)

        self.txt_cedula = ctk.CTkEntry(self.frame_form, placeholder_text="Cédula (Sólo números)", fg_color="#2D2924")
        self.txt_cedula.pack(pady=3, fill="x", padx=10)

        self.txt_nombre = ctk.CTkEntry(self.frame_form, placeholder_text="Nombre", fg_color="#2D2924")
        self.txt_nombre.pack(pady=3, fill="x", padx=10)

        self.txt_apellido = ctk.CTkEntry(self.frame_form, placeholder_text="Apellido", fg_color="#2D2924")
        self.txt_apellido.pack(pady=3, fill="x", padx=10)

        self.txt_correo = ctk.CTkEntry(self.frame_form, placeholder_text="Correo Electrónico", fg_color="#2D2924")
        self.txt_correo.pack(pady=3, fill="x", padx=10)

        self.txt_telefono = ctk.CTkEntry(self.frame_form, placeholder_text="Teléfono", fg_color="#2D2924")
        self.txt_telefono.pack(pady=3, fill="x", padx=10)

        self.txt_usuario = ctk.CTkEntry(self.frame_form, placeholder_text="Usuario de Acceso", fg_color="#2D2924")
        self.txt_usuario.pack(pady=3, fill="x", padx=10)

        # Campo Contraseña con Botón de Ojito
        f_pass = ctk.CTkFrame(self.frame_form, fg_color="transparent")
        f_pass.pack(fill="x", padx=10, pady=3)

        self.txt_clave = ctk.CTkEntry(f_pass, placeholder_text="Contraseña (Vacío = Mantener)", show="*", fg_color="#2D2924")
        self.txt_clave.pack(side="left", fill="x", expand=True, padx=(0, 4))

        self.btn_ojito = ctk.CTkButton(f_pass, text="👁️", width=36, fg_color="#3E342B", hover_color="#524539", command=self.toggle_ver_clave)
        self.btn_ojito.pack(side="right")

        self.combo_rol = ctk.CTkComboBox(self.frame_form, values=["recepcionista", "admin"])
        self.combo_rol.set("recepcionista")
        self.combo_rol.pack(pady=8, fill="x", padx=10)

        self.btn_guardar = ctk.CTkButton(self.frame_form, text="Registrar Usuario", fg_color="#D4A343", hover_color="#B8892E", text_color="#191715", font=("Arial", 11, "bold"), command=self.guardar_usuario)
        self.btn_guardar.pack(pady=10, fill="x", padx=10)

        self.btn_cancelar = ctk.CTkButton(self.frame_form, text="❌ Cancelar Edición", fg_color="#78281F", hover_color="#943126", command=self.limpiar_campos)

        self.frame_lista = ctk.CTkScrollableFrame(self.container_pestanas, fg_color="#23201C", label_text="Personal Registrado")
        self.frame_lista.grid(row=0, column=1, padx=10, pady=10, sticky="nsew")

        self.cargar_usuarios()

    def toggle_ver_clave(self):
        """Alterna entre mostrar y ocultar los caracteres de la contraseña."""
        self.ver_clave_activa = not self.ver_clave_activa
        if self.ver_clave_activa:
            self.txt_clave.configure(show="")
            self.btn_ojito.configure(text="🔒")
        else:
            self.txt_clave.configure(show="*")
            self.btn_ojito.configure(text="👁️")

    def guardar_usuario(self):
        cedula, nombre, apellido, correo = self.txt_cedula.get().strip(), self.txt_nombre.get().strip(), self.txt_apellido.get().strip(), self.txt_correo.get().strip()
        telefono, usuario, clave, rol = self.txt_telefono.get().strip(), self.txt_usuario.get().strip(), self.txt_clave.get().strip(), self.combo_rol.get()

        if not all([cedula, nombre, apellido, correo, telefono, usuario]):
            messagebox.showwarning("Atención", "Complete todos los campos obligatorios.")
            return

        admin_act = self.controller.usuario_actual["usuario"]
        rol_act = self.controller.usuario_actual["rol"]

        if self.usuario_id_edicion:
            clave_hash = hash_clave(clave) if clave else None
            if actualizar_usuario(self.usuario_id_edicion, cedula, nombre, apellido, correo, telefono, usuario, clave_hash, rol):
                registrar_auditoria(admin_act, rol_act, "MODIFICAR_USUARIO", f"Se actualizaron los datos del usuario '{usuario}' (C.I: {cedula}, Rol: {rol}).")
                messagebox.showinfo("Éxito", "Usuario actualizado correctamente.")
                self.limpiar_campos()
                self.cargar_usuarios()
            else:
                messagebox.showerror("Error", "Cédula o Usuario ya en uso.")
        else:
            if not clave:
                messagebox.showwarning("Atención", "Ingrese una contraseña para el nuevo usuario.")
                return
            conn = obtener_conexion()
            cursor = conn.cursor()
            try:
                cursor.execute("""
                    INSERT INTO usuarios (cedula, nombre, apellido, correo, telefono, usuario, clave, rol)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?);
                """, (cedula, nombre, apellido, correo, telefono, usuario, hash_clave(clave), rol))
                conn.commit()
                registrar_auditoria(admin_act, rol_act, "CREAR_USUARIO", f"Se creó el nuevo usuario '{usuario}' (C.I: {cedula}, Rol: {rol}).")
                messagebox.showinfo("Éxito", "Usuario registrado.")
                self.limpiar_campos()
                self.cargar_usuarios()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo registrar: {e}")
            finally:
                conn.close()

    def seleccionar_para_editar(self, user: dict):
        self.usuario_id_edicion = user["id"]
        self.txt_cedula.delete(0, 'end'); self.txt_cedula.insert(0, user["cedula"])
        self.txt_nombre.delete(0, 'end'); self.txt_nombre.insert(0, user["nombre"])
        self.txt_apellido.delete(0, 'end'); self.txt_apellido.insert(0, user["apellido"])
        self.txt_correo.delete(0, 'end'); self.txt_correo.insert(0, user["correo"])
        self.txt_telefono.delete(0, 'end'); self.txt_telefono.insert(0, user["telefono"])
        self.txt_usuario.delete(0, 'end'); self.txt_usuario.insert(0, user["usuario"])
        self.txt_clave.delete(0, 'end')
        self.combo_rol.set(user["rol"])

        self.lbl_titulo_form.configure(text=f"Modificando ID #{user['id']}")
        self.btn_guardar.configure(text="💾 Actualizar Datos", fg_color="#3E342B", text_color="#F4EFE6")
        self.btn_cancelar.pack(pady=4, fill="x", padx=10)

    def limpiar_campos(self):
        self.usuario_id_edicion = None
        for txt in [self.txt_cedula, self.txt_nombre, self.txt_apellido, self.txt_correo, self.txt_telefono, self.txt_usuario, self.txt_clave]:
            txt.delete(0, 'end')
        self.combo_rol.set("recepcionista")
        self.lbl_titulo_form.configure(text="Gestión de Personal")
        self.btn_guardar.configure(text="Registrar Usuario", fg_color="#D4A343", text_color="#191715")
        self.btn_cancelar.pack_forget()

    def cargar_usuarios(self):
        for w in self.frame_lista.winfo_children(): w.destroy()
        for u in obtener_todos_usuarios():
            card = ctk.CTkFrame(self.frame_lista, fg_color="#2D2924")
            card.pack(fill="x", pady=3, padx=5)
            lbl_t = f"C.I: {u['cedula']} | {u['nombre']} {u['apellido']}\nUser: {u['usuario']} | Rol: {u['rol'].upper()} | Tel: {u['telefono']}"
            ctk.CTkLabel(card, text=lbl_t, font=("Arial", 11), justify="left", text_color="#F4EFE6").pack(side="left", padx=10, pady=6)
            ctk.CTkButton(card, text="✏️ Editar", fg_color="#3E342B", hover_color="#524539", width=70, command=lambda usr=u: self.seleccionar_para_editar(usr)).pack(side="right", padx=10)

    # =========================================================================
    # PESTAÑA 2: AUDITORÍA CON FILTRO POR USUARIOS
    # =========================================================================
    def construir_pestana_auditoria(self):
        frame_filtros = ctk.CTkFrame(self.container_pestanas, fg_color="#23201C", corner_radius=12, border_width=1, border_color="#36322C")
        frame_filtros.pack(fill="x", padx=10, pady=5)

        ctk.CTkLabel(frame_filtros, text="Desde:", font=("Arial", 10, "bold"), text_color="#D4A343").pack(side="left", padx=4)
        self.cal_aud_ini = DateEntry(frame_filtros, date_pattern='yyyy-mm-dd')
        self.cal_aud_ini.pack(side="left", padx=2)

        ctk.CTkLabel(frame_filtros, text="Hasta:", font=("Arial", 10, "bold"), text_color="#D4A343").pack(side="left", padx=4)
        self.cal_aud_fin = DateEntry(frame_filtros, date_pattern='yyyy-mm-dd')
        self.cal_aud_fin.pack(side="left", padx=2)

        # Filtro de Usuario Dinámico
        ctk.CTkLabel(frame_filtros, text="Usuario:", font=("Arial", 10, "bold"), text_color="#D4A343").pack(side="left", padx=4)
        usuarios_lista = ["TODOS"] + [u["usuario"] for u in obtener_todos_usuarios()]
        self.combo_filtro_usuario = ctk.CTkComboBox(frame_filtros, values=usuarios_lista, width=110)
        self.combo_filtro_usuario.set("TODOS")
        self.combo_filtro_usuario.pack(side="left", padx=2)

        # Filtro de Acción
        ctk.CTkLabel(frame_filtros, text="Acción:", font=("Arial", 10, "bold"), text_color="#D4A343").pack(side="left", padx=4)
        self.combo_filtro_accion = ctk.CTkComboBox(
            frame_filtros, width=130,
            values=["TODAS", "LOGIN", "LOGIN_FALLIDO", "CHECK_IN", "CHECK_OUT", "EXTENSION", "VENTA_MINIBAR", "CAMBIO_TASA", "CREAR_USUARIO", "MODIFICAR_USUARIO", "BLOQUEAR_CLIENTE", "DESBLOQUEAR_CLIENTE", "CIERRE_TURNO", "CREAR_RESPALDO", "RESETEO_BD"]
        )
        self.combo_filtro_accion.set("TODAS")
        self.combo_filtro_accion.pack(side="left", padx=2)

        ctk.CTkButton(frame_filtros, text="🔍 Filtrar", fg_color="#3E342B", hover_color="#524539", width=70, command=self.cargar_logs_pantalla).pack(side="left", padx=6)
        ctk.CTkButton(frame_filtros, text="📄 PDF", fg_color="#78281F", hover_color="#943126", width=70, command=self.exportar_auditoria_pdf).pack(side="right", padx=3)
        ctk.CTkButton(frame_filtros, text="📊 Excel", fg_color="#1E5F38", hover_color="#2E7D32", width=70, command=self.exportar_auditoria_excel).pack(side="right", padx=3)

        self.scroll_logs = ctk.CTkScrollableFrame(self.container_pestanas, fg_color="#23201C", label_text="Bitácora de Eventos y Trazabilidad")
        self.scroll_logs.pack(fill="both", expand=True, padx=10, pady=8)

        self.logs_cargados = []
        self.cargar_logs_pantalla()

    def cargar_logs_pantalla(self):
        for w in self.scroll_logs.winfo_children(): w.destroy()

        f1 = self.cal_aud_ini.get_date().strftime("%Y-%m-%d")
        f2 = self.cal_aud_fin.get_date().strftime("%Y-%m-%d")
        acc = self.combo_filtro_accion.get()
        usr = self.combo_filtro_usuario.get()

        self.logs_cargados = obtener_logs_auditoria(f1, f2, acc, usr)

        if not self.logs_cargados:
            ctk.CTkLabel(self.scroll_logs, text="No hay registros de auditoría para los filtros aplicados.", font=("Arial", 12, "italic"), text_color="#A89F91").pack(pady=30)
            return

        color_badges = {
            "LOGIN": "#1E5F38", "LOGIN_FALLIDO": "#78281F", "CHECK_IN": "#D4A343",
            "CHECK_OUT": "#A05A18", "EXTENSION": "#632D73", "VENTA_MINIBAR": "#2E86C1",
            "CIERRE_TURNO": "#8E44AD", "RESETEO_BD": "#C0392B", "BLOQUEAR_CLIENTE": "#922B21"
        }

        for log in self.logs_cargados:
            bg_badge = color_badges.get(log["accion"], "#3E342B")
            card = ctk.CTkFrame(self.scroll_logs, fg_color="#2D2924")
            card.pack(fill="x", pady=2, padx=5)

            lbl_badge = ctk.CTkLabel(card, text=f" {log['accion']} ", fg_color=bg_badge, corner_radius=6, font=("Arial", 9, "bold"), text_color="#FFFFFF")
            lbl_badge.pack(side="left", padx=8, pady=6)

            txt = f"{log['fecha']} | Usuario: {log['usuario']} ({log['rol'].upper()}) -> {log['descripcion']}"
            ctk.CTkLabel(card, text=txt, font=("Arial", 10), text_color="#F4EFE6", justify="left").pack(side="left", padx=5)

    def exportar_auditoria_pdf(self):
        if not self.logs_cargados:
            messagebox.showwarning("Atención", "No hay datos para exportar.")
            return

        pdf = PDFReporteAuditoria()
        pdf.add_page()
        pdf.set_font("Arial", "", 8)

        for l in self.logs_cargados:
            pdf.cell(32, 6, str(l["fecha"]), 1, 0, "C")
            pdf.cell(24, 6, str(l["usuario"])[:15], 1, 0, "C")
            pdf.cell(20, 6, str(l["rol"]).upper()[:12], 1, 0, "C")
            pdf.cell(35, 6, str(l["accion"])[:20], 1, 0, "C")
            pdf.cell(166, 6, str(l["descripcion"])[:110], 1, 0, "L")
            pdf.ln()

        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        export_dir = os.path.join(base_dir, "reportes_exportados")
        os.makedirs(export_dir, exist_ok=True)

        ruta = os.path.join(export_dir, f"auditoria_personal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta)
        os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_auditoria_excel(self):
        if not self.logs_cargados:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Auditoría de Personal"
        ws.append(["Fecha / Hora", "Usuario", "Rol", "Tipo de Acción", "Descripción Detallada"])

        for l in self.logs_cargados:
            ws.append([l["fecha"], l["usuario"], l["rol"], l["accion"], l["descripcion"]])

        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        export_dir = os.path.join(base_dir, "reportes_exportados")
        os.makedirs(export_dir, exist_ok=True)

        ruta = os.path.join(export_dir, f"auditoria_personal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(ruta)
        os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')