"""
===============================================================================
SISTEMA DE GESTIÓN DE MOTELES - MOTEL PRO v2.0
Módulo: modules/reportes.py (CENTRO DE REPORTES: SAIME, P2P, CPP Y HOUSEKEEPING)
Diseño: Dark Luxury Boutique & Enterprise Auditing Suite
===============================================================================
"""

import os
import openpyxl
import customtkinter as ctk
from tkinter import messagebox
from tkcalendar import DateEntry
from fpdf import FPDF
from datetime import datetime

# Importaciones desde la Capa de Datos (Modelo SQLite)
from database.db_manager import (
    obtener_historial_saime, obtener_reporte_p2p, obtener_reporte_utilidad_inventario,
    obtener_historial_limpieza, obtener_tasa_bcv, formatear_bs, formatear_usd,
    obtener_ruta_recurso
)


# =============================================================================
# 1. GENERADORES DE REPORTES PDF FORMALES
# =============================================================================
class PDFPlanillaSAIME(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=10)

    def header(self):
        logo_h = obtener_ruta_recurso("logo_hotel.png")
        logo_s = obtener_ruta_recurso("logo_saime.png")
        if os.path.exists(logo_h): self.image(logo_h, x=10, y=6, w=22)
        if os.path.exists(logo_s): self.image(logo_s, x=256, y=6, w=30)

        self.set_font("Helvetica", "B", 9)
        self.cell(0, 4, "REPÚBLICA BOLIVARIANA DE VENEZUELA", 0, 1, "C")
        self.cell(0, 4, "MINISTERIO DEL PODER POPULAR PARA RELACIONES INTERIORES, JUSTICIA Y PAZ", 0, 1, "C")
        self.set_font("Helvetica", "B", 10)
        self.cell(0, 4.5, "DIRECCIÓN GENERAL DE MIGRACIÓN Y EXTRANJERÍA - SAIME", 0, 1, "C")
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN | RIF: J-30250227-6", 0, 1, "C")
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 4, "REGISTRO DIARIO DE HUÉSPEDES Y CONTROL DE EXTRANJERÍA", 0, 1, "C")
        self.set_y(32)

        self.set_fill_color(33, 30, 27)
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 8)
        self.cell(12, 7, "N° PER", 1, 0, "C", fill=True)
        self.cell(22, 7, "FECHA", 1, 0, "C", fill=True)
        self.cell(65, 7, "NOMBRE Y APELLIDO", 1, 0, "C", fill=True)
        self.cell(12, 7, "EDAD", 1, 0, "C", fill=True)
        self.cell(22, 7, "ESTADO CIVIL", 1, 0, "C", fill=True)
        self.cell(10, 7, "NAC", 1, 0, "C", fill=True)
        self.cell(45, 7, "PROCEDENCIA", 1, 0, "C", fill=True)
        self.cell(27, 7, "CÉDULA", 1, 0, "C", fill=True)
        self.cell(62, 7, "DESTINO", 1, 0, "C", fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"Planilla Oficial para Inspección Policial / SAIME - Página {self.page_no()}", 0, 0, "C")


class PDFReporteP2P(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)

    def header(self):
        logo = obtener_ruta_recurso("logo_hotel.png")
        if os.path.exists(logo): self.image(logo, x=10, y=6, w=20)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN", 0, 1, "C")
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 4.5, "RIF: J-30250227-6 | REPORTE DE AUDITORÍA DE TRANSACCIONES PAGO MÓVIL (P2P)", 0, 1, "C")
        self.set_y(26)

        self.set_fill_color(33, 30, 27)
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 8)
        self.cell(30, 7, "FECHA / HORA", 1, 0, "C", fill=True)
        self.cell(28, 7, "HABITACIÓN", 1, 0, "C", fill=True)
        self.cell(38, 7, "CONCEPTO", 1, 0, "C", fill=True)
        self.cell(50, 7, "CLIENTE / TITULAR", 1, 0, "C", fill=True)
        self.cell(24, 7, "C.I. EMISOR", 1, 0, "C", fill=True)
        self.cell(28, 7, "TELÉFONO", 1, 0, "C", fill=True)
        self.cell(22, 7, "REF (P2P)", 1, 0, "C", fill=True)
        self.cell(32, 7, "MONTO (BS)", 1, 0, "C", fill=True)
        self.cell(25, 7, "MONTO ($)", 1, 0, "C", fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"Auditoría Contable Bancaria P2P - Página {self.page_no()}", 0, 0, "C")


class PDFReporteUtilidadCPP(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)

    def header(self):
        logo = obtener_ruta_recurso("logo_hotel.png")
        if os.path.exists(logo): self.image(logo, x=10, y=6, w=20)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN", 0, 1, "C")
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 4.5, "RIF: J-30250227-6 | ESTADO DE RESULTADOS: UTILIDAD Y GANANCIAS MINI-BAR (CPP)", 0, 1, "C")
        self.set_y(26)

        self.set_fill_color(33, 30, 27)
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 7.5)
        self.cell(42, 7, "PRODUCTO", 1, 0, "L", fill=True)
        self.cell(14, 7, "CANT", 1, 0, "C", fill=True)
        self.cell(24, 7, "COSTO ($)", 1, 0, "R", fill=True)
        self.cell(24, 7, "VENTA ($)", 1, 0, "R", fill=True)
        self.cell(28, 7, "TOT. COSTO ($)", 1, 0, "R", fill=True)
        self.cell(28, 7, "TOT. VENTA ($)", 1, 0, "R", fill=True)
        self.cell(38, 7, "TOT. VENTA (BS)", 1, 0, "R", fill=True)
        self.cell(28, 7, "UTILIDAD ($)", 1, 0, "R", fill=True)
        self.cell(34, 7, "UTILIDAD (BS)", 1, 0, "R", fill=True)
        self.cell(17, 7, "MARGEN", 1, 0, "C", fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"Auditoría Contable de Rentabilidad - Página {self.page_no()}", 0, 0, "C")


class PDFReporteLimpieza(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)

    def header(self):
        logo = obtener_ruta_recurso("logo_hotel.png")
        if os.path.exists(logo): self.image(logo, x=10, y=6, w=20)
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN", 0, 1, "C")
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 4.5, "RIF: J-30250227-6 | REGISTRO DE LIMPIEZAS Y CONTROL DE CAMARERAS (SLA)", 0, 1, "C")
        self.set_y(26)

        self.set_fill_color(33, 30, 27)
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 8)
        self.cell(32, 7, "FECHA / HORA LIMPIA", 1, 0, "C", fill=True)
        self.cell(25, 7, "HABITACIÓN", 1, 0, "C", fill=True)
        self.cell(60, 7, "CAMARERA RESPONSABLE", 1, 0, "L", fill=True)
        self.cell(25, 7, "TURNO", 1, 0, "C", fill=True)
        self.cell(50, 7, "HORA QUE PASÓ A SUCIA", 1, 0, "C", fill=True)
        self.cell(45, 7, "TIEMPO EN SUCIA (MIN)", 1, 0, "C", fill=True)
        self.cell(40, 7, "ESTADO CALIDAD", 1, 0, "C", fill=True)
        self.ln()
        self.set_text_color(0, 0, 0)

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, f"Auditoría de Camareras y Housekeeping - Página {self.page_no()}", 0, 0, "C")


# =============================================================================
# 2. VISTA PRINCIPAL CON DISEÑO DARK LUXURY Y BUSCADOR INTEGRADO
# =============================================================================
class FrameReportes(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, fg_color="#100F0D")
        self.controller = controller

        # Paleta Corporativa Dark Luxury
        self.C_BG = "#100F0D"
        self.C_PANEL = "#161513"
        self.C_CARD = "#1E1C19"
        self.C_INPUT = "#26231F"
        self.C_BORDER = "#2E2A25"
        self.C_GOLD = "#D4AF37"
        self.C_GOLD_HOVER = "#B89228"
        self.C_GREEN = "#2ECC71"
        self.C_RED = "#E74C3C"
        self.C_BLUE = "#3498DB"
        self.C_TEXT_MAIN = "#F5EFEB"
        self.C_TEXT_MUTED = "#8E8880"

        # Caché de datos en memoria
        self.datos_saime = []
        self.datos_p2p = []
        self.datos_cpp = []
        self.datos_limpieza = []

        # =====================================================================
        # TARJETA SUPERIOR: NAVEGACIÓN Y FILTROS TEMPORALES
        # =====================================================================
        self.frame_top = ctk.CTkFrame(
            self, fg_color=self.C_PANEL, corner_radius=14,
            border_width=1, border_color=self.C_BORDER
        )
        self.frame_top.pack(fill="x", padx=16, pady=(16, 8))

        # Fila 1: Título y Segmented Button
        f_row1 = ctk.CTkFrame(self.frame_top, fg_color="transparent")
        f_row1.pack(fill="x", padx=18, pady=(14, 8))

        f_title_box = ctk.CTkFrame(f_row1, fg_color="transparent")
        f_title_box.pack(side="left")

        ctk.CTkLabel(
            f_title_box, text="CENTRO DE REPORTES Y AUDITORÍA",
            font=("Montserrat", 14, "bold"), text_color=self.C_GOLD
        ).pack(anchor="w")

        ctk.CTkLabel(
            f_title_box, text="Libros oficiales SAIME, transacciones P2P, rentabilidad CPP y control SLA",
            font=("Arial", 9), text_color=self.C_TEXT_MUTED
        ).pack(anchor="w")

        self.tab_selector = ctk.CTkSegmentedButton(
            f_row1,
            values=["📑 Extranjería (SAIME)", "💳 Auditoría P2P", "📈 Utilidad Mini-Bar", "🧹 Control Camareras"],
            selected_color=self.C_GOLD, selected_hover_color=self.C_GOLD_HOVER,
            unselected_color=self.C_CARD, unselected_hover_color="#DF6721",
            text_color="#F7F6F3", font=("Arial", 11, "bold"),
            command=lambda v: self.cargar_datos()
        )
        self.tab_selector.set("📑 Extranjería (SAIME)")
        self.tab_selector.pack(side="right")

        # Fila 2: Filtros de Fecha, Buscador y Botones de Exportación
        f_row2 = ctk.CTkFrame(self.frame_top, fg_color="transparent")
        f_row2.pack(fill="x", padx=18, pady=(0, 14))

        f_fechas = ctk.CTkFrame(f_row2, fg_color="transparent")
        f_fechas.pack(side="left")

        ctk.CTkLabel(f_fechas, text="Desde:", font=("Arial", 10, "bold"), text_color=self.C_TEXT_MUTED).pack(side="left", padx=4)
        self.cal_inicio = DateEntry(
            f_fechas, date_pattern='yyyy-mm-dd',
            background="#211E1B", foreground="white", headersbackground="#D4AF37"
        )
        self.cal_inicio.set_date(datetime.now().date().replace(day=1))
        self.cal_inicio.pack(side="left", padx=4)

        ctk.CTkLabel(f_fechas, text="Hasta:", font=("Arial", 10, "bold"), text_color=self.C_TEXT_MUTED).pack(side="left", padx=6)
        self.cal_fin = DateEntry(
            f_fechas, date_pattern='yyyy-mm-dd',
            background="#211E1B", foreground="white", headersbackground="#D4AF37"
        )
        self.cal_fin.pack(side="left", padx=4)

        ctk.CTkButton(
            f_fechas, text="🔍 Consultar", fg_color=self.C_CARD, hover_color="#2A2621",
            text_color=self.C_GOLD, width=100, height=32, corner_radius=8,
            border_width=1, border_color=self.C_BORDER, font=("Arial", 11, "bold"),
            command=self.cargar_datos
        ).pack(side="left", padx=10)

        # Buscador interactivo
        self.entry_search = ctk.CTkEntry(
            f_row2, placeholder_text="🔍 Filtrar registros en pantalla...",
            width=260, height=32, fg_color=self.C_INPUT, border_color=self.C_BORDER,
            text_color=self.C_TEXT_MAIN, font=("Arial", 10)
        )
        self.entry_search.pack(side="left", padx=15)
        self.entry_search.bind("<KeyRelease>", lambda e: self.filtrar_pantalla())

        # Botones de Exportación
        f_acciones = ctk.CTkFrame(f_row2, fg_color="transparent")
        f_acciones.pack(side="right")

        self.btn_export_pdf = ctk.CTkButton(
            f_acciones, text="📄 Exportar PDF", fg_color="#3A1714", hover_color=self.C_RED,
            text_color="#FADBD8", width=120, height=32, corner_radius=8,
            border_width=1, border_color="#5C201A", font=("Arial", 11, "bold"),
            command=self.exportar_pdf_actual
        )
        self.btn_export_pdf.pack(side="left", padx=4)

        self.btn_export_excel = ctk.CTkButton(
            f_acciones, text="📊 Exportar Excel", fg_color="#1E5F38", hover_color="#2E7D32",
            text_color="#F4EFE6", width=120, height=32, corner_radius=8,
            font=("Arial", 11, "bold"), command=self.exportar_excel_actual
        )
        self.btn_export_excel.pack(side="left", padx=4)

        # =====================================================================
        # TARJETA RESUMEN FLOTANTE (KPIS POR PESTAÑA)
        # =====================================================================
        self.card_kpi_resumen = ctk.CTkFrame(
            self, fg_color=self.C_PANEL, corner_radius=10,
            border_width=1, border_color=self.C_BORDER
        )
        self.card_kpi_resumen.pack(fill="x", padx=16, pady=(0, 8))

        self.lbl_resumen_kpi = ctk.CTkLabel(
            self.card_kpi_resumen, text="Cargando estadísticas...",
            font=("Montserrat", 11, "bold"), text_color=self.C_GOLD
        )
        self.lbl_resumen_kpi.pack(pady=8, padx=16, anchor="w")

        # =====================================================================
        # CONTENEDOR SCROLLABLE DE REGISTROS
        # =====================================================================
        self.scroll_data = ctk.CTkScrollableFrame(self, fg_color=self.C_BG)
        self.scroll_data.pack(fill="both", expand=True, padx=16, pady=(0, 16))

        self.cargar_datos()

    def obtener_ruta_exportacion(self, file_name: str) -> str:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        export_dir = os.path.join(base_dir, "reportes_exportados")
        os.makedirs(export_dir, exist_ok=True)
        return os.path.join(export_dir, file_name)

    # =========================================================================
    # CARGA Y DESPLIEGUE DINÁMICO DE DATOS
    # =========================================================================
    def cargar_datos(self):
        f1 = self.cal_inicio.get_date().strftime("%Y-%m-%d")
        f2 = self.cal_fin.get_date().strftime("%Y-%m-%d")
        tipo = self.tab_selector.get()
        tasa = obtener_tasa_bcv()

        if "SAIME" in tipo:
            self.datos_saime = obtener_historial_saime(f1, f2)
            tot_huespedes = 0
            for e in self.datos_saime:
                tot_huespedes += 1  # Titular
                if e.get("ac_nombre") and e.get("ac_cedula") and e["ac_nombre"].strip():
                    tot_huespedes += 1  # Acompañante
            self.lbl_resumen_kpi.configure(
                text=f"📋 Historial SAIME: {len(self.datos_saime)} Habitaciones Ocupadas  •  {tot_huespedes} Personas Registradas en el Período",
                text_color=self.C_GOLD
            )

        elif "P2P" in tipo:
            self.datos_p2p = obtener_reporte_p2p(f1, f2)
            tot_usd = sum(p["monto_usd"] for p in self.datos_p2p)
            tot_bs = sum(p["monto_bs"] for p in self.datos_p2p)
            self.lbl_resumen_kpi.configure(
                text=f"📱 Conciliación Pago Móvil: {len(self.datos_p2p)} Operaciones  •  Total: ${tot_usd:,.2f} ({formatear_bs(tot_bs)})",
                text_color=self.C_BLUE
            )

        elif "Utilidad" in tipo:
            self.datos_cpp = obtener_reporte_utilidad_inventario(f1, f2)
            tot_ing = sum(d["ingreso_total_usd"] for d in self.datos_cpp)
            tot_cos = sum(d["costo_total_usd"] for d in self.datos_cpp)
            tot_uti = sum(d["utilidad_neta_usd"] for d in self.datos_cpp)
            margen_g = (tot_uti / tot_ing * 100) if tot_ing > 0 else 0.0
            self.lbl_resumen_kpi.configure(
                text=f"💰 Utilidad Mini-Bar: Ventas ${tot_ing:,.2f}  |  Costo CPP ${tot_cos:,.2f}  |  GANANCIA NETA: +${tot_uti:,.2f} ({formatear_bs(tot_uti * tasa)})  [Margen: {margen_g:.1f}%]",
                text_color=self.C_GREEN
            )

        else:
            self.datos_limpieza = obtener_historial_limpieza(f1, f2)
            prom_sla = (sum(l["duracion_minutos"] for l in self.datos_limpieza) / len(self.datos_limpieza)) if self.datos_limpieza else 0
            self.lbl_resumen_kpi.configure(
                text=f"🧹 Auditoría de Housekeeping: {len(self.datos_limpieza)} Limpiezas Registradas  •  Tiempo Promedio de Rotación: {prom_sla:.1f} minutos",
                text_color=self.C_GOLD
            )

        self.filtrar_pantalla()

    def filtrar_pantalla(self):
        for w in self.scroll_data.winfo_children():
            w.destroy()

        query = self.entry_search.get().strip().upper()
        tipo = self.tab_selector.get()

        if "SAIME" in tipo:
            if not self.datos_saime:
                self._render_vacio("No hay registros de extranjería / SAIME en el rango de fechas seleccionado.")
                return

            nro = 1
            for est in self.datos_saime:
                tit_nom = f"{est['nombre']} {est['apellido']}".upper()
                tit_ci = est['cedula']
                ac_nom = f"{est['ac_nombre']} {est['ac_apellido']}".upper() if est.get("ac_nombre") else ""
                ac_ci = est.get("ac_cedula", "")
                hab = est['hab_codigo'].upper()

                if query and not any(query in campo for campo in [tit_nom, tit_ci, ac_nom, ac_ci, hab]):
                    nro += (2 if ac_nom.strip() else 1)
                    continue

                fecha_f = datetime.strptime(est["fecha_entrada"], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %I:%M %p")

                # Fila Titular
                card_t = ctk.CTkFrame(self.scroll_data, fg_color=self.C_CARD, corner_radius=8, border_width=1, border_color=self.C_BORDER)
                card_t.pack(fill="x", pady=2, padx=4)

                f_t_izq = ctk.CTkFrame(card_t, fg_color="transparent")
                f_t_izq.pack(side="left", padx=12, pady=8)

                ctk.CTkLabel(f_t_izq, text=f"#{nro:04d}  •  {tit_nom}  (C.I: {tit_ci})", font=("Montserrat", 11, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="w")
                ctk.CTkLabel(f_t_izq, text=f"Fecha: {fecha_f}  |  Hab: {hab}  |  Edad: {est['edad']} años  |  Ruta: {est['procedencia']} ➔ {est['destino']}", font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="w")

                badge_t = ctk.CTkFrame(card_t, fg_color="#211E1B", corner_radius=5, border_width=1, border_color=self.C_GOLD)
                badge_t.pack(side="right", padx=12, pady=8)
                ctk.CTkLabel(badge_t, text="TITULAR", font=("Arial", 8, "bold"), text_color=self.C_GOLD).pack(padx=6, pady=2)
                nro += 1

                # Fila Acompañante
                if est.get("ac_nombre") and est.get("ac_cedula") and est["ac_nombre"].strip():
                    card_a = ctk.CTkFrame(self.scroll_data, fg_color="#181614", corner_radius=8, border_width=1, border_color="#26231F")
                    card_a.pack(fill="x", pady=2, padx=4)

                    f_a_izq = ctk.CTkFrame(card_a, fg_color="transparent")
                    f_a_izq.pack(side="left", padx=12, pady=8)

                    ctk.CTkLabel(f_a_izq, text=f"#{nro:04d}  •  {ac_nom}  (C.I: {ac_ci})", font=("Montserrat", 11, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="w")
                    ctk.CTkLabel(f_a_izq, text=f"Fecha: {fecha_f}  |  Hab: {hab}  |  Edad: {est['ac_edad']} años  |  Ruta: {est['ac_procedencia']} ➔ {est['ac_destino']}", font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="w")

                    badge_a = ctk.CTkFrame(card_a, fg_color="#161513", corner_radius=5, border_width=1, border_color=self.C_TEXT_MUTED)
                    badge_a.pack(side="right", padx=12, pady=8)
                    ctk.CTkLabel(badge_a, text="ACOMPAÑANTE", font=("Arial", 8, "bold"), text_color=self.C_TEXT_MUTED).pack(padx=6, pady=2)
                    nro += 1

        elif "P2P" in tipo:
            if not self.datos_p2p:
                self._render_vacio("No se registraron cobros con Pago Móvil en este período.")
                return

            for p in self.datos_p2p:
                cli = p['cliente_nombre'].upper()
                ref = str(p.get('p2p_referencia', '')).upper()
                ci = str(p.get('p2p_ci', '')).upper()
                hab = p['hab_codigo'].upper()
                con = p['concepto'].upper()

                if query and not any(query in campo for campo in [cli, ref, ci, hab, con]):
                    continue

                card = ctk.CTkFrame(self.scroll_data, fg_color=self.C_CARD, corner_radius=8, border_width=1, border_color=self.C_BORDER)
                card.pack(fill="x", pady=3, padx=4)

                f_izq = ctk.CTkFrame(card, fg_color="transparent")
                f_izq.pack(side="left", padx=12, pady=8)

                ctk.CTkLabel(f_izq, text=f"Hab {p['hab_codigo']}  •  {p['concepto']}  •  {p['cliente_nombre']}", font=("Montserrat", 11, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="w")
                ctk.CTkLabel(f_izq, text=f"Fecha: {p['fecha']}  |  Ref: {p['p2p_referencia']}  |  C.I: {p['p2p_ci']}  |  Tel: {p['p2p_telefono']}  |  Cajero: {p['recepcionista']}", font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="w")

                f_der = ctk.CTkFrame(card, fg_color="transparent")
                f_der.pack(side="right", padx=12, pady=8)

                ctk.CTkLabel(f_der, text=f"${p['monto_usd']:.2f}", font=("Montserrat", 13, "bold"), text_color=self.C_GOLD).pack(anchor="e")
                ctk.CTkLabel(f_der, text=formatear_bs(p['monto_bs']), font=("Arial", 9, "bold"), text_color=self.C_GREEN).pack(anchor="e")

        elif "Utilidad" in tipo:
            if not self.datos_cpp:
                self._render_vacio("No hay ventas registradas de mini-bar para calcular utilidad.")
                return

            for d in self.datos_cpp:
                prod = d['producto'].upper()
                if query and query not in prod:
                    continue

                card = ctk.CTkFrame(self.scroll_data, fg_color=self.C_CARD, corner_radius=8, border_width=1, border_color=self.C_BORDER)
                card.pack(fill="x", pady=3, padx=4)

                f_izq = ctk.CTkFrame(card, fg_color="transparent")
                f_izq.pack(side="left", padx=12, pady=8)

                ctk.CTkLabel(f_izq, text=f"🛒 {d['producto']}", font=("Montserrat", 11, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="w")
                ctk.CTkLabel(f_izq, text=f"Vendidos: {d['unidades_vendidas']} un.  |  Venta Total: ${d['ingreso_total_usd']:.2f}  |  Costo Total (CPP): ${d['costo_total_usd']:.2f}", font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="w")

                f_der = ctk.CTkFrame(card, fg_color="transparent")
                f_der.pack(side="right", padx=12, pady=8)

                ctk.CTkLabel(f_der, text=f"+${d['utilidad_neta_usd']:.2f} ({d['margen_pct']:.1f}%)", font=("Montserrat", 12, "bold"), text_color=self.C_GREEN).pack(anchor="e")
                ctk.CTkLabel(f_der, text=formatear_bs(d['utilidad_neta_bs']), font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="e")

        else:
            if not self.datos_limpieza:
                self._render_vacio("No hay registros de limpieza en el período seleccionado.")
                return

            for lim in self.datos_limpieza:
                hab = lim['hab_codigo'].upper()
                cam = lim['camarera'].upper()

                if query and not any(query in campo for campo in [hab, cam]):
                    continue

                minutos = lim["duracion_minutos"]
                es_lenta = minutos > 45
                color_calidad = self.C_RED if es_lenta else self.C_GREEN
                txt_calidad = "DEMORADO (>45 min)" if es_lenta else "ÓPTIMO (SLA OK)"

                card = ctk.CTkFrame(self.scroll_data, fg_color=self.C_CARD, corner_radius=8, border_width=1, border_color=self.C_BORDER)
                card.pack(fill="x", pady=3, padx=4)

                f_izq = ctk.CTkFrame(card, fg_color="transparent")
                f_izq.pack(side="left", padx=12, pady=8)

                ctk.CTkLabel(f_izq, text=f"Habitación {lim['hab_codigo']}  •  Camarera: {lim['camarera']}", font=("Montserrat", 11, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="w")
                ctk.CTkLabel(f_izq, text=f"Liberada: {lim['fecha_limpia']}  |  Pasó a Sucia: {lim['fecha_sucia'] or 'N/A'}  |  Turno: {lim['turno']}", font=("Arial", 9), text_color=self.C_TEXT_MUTED).pack(anchor="w")

                f_der = ctk.CTkFrame(card, fg_color="transparent")
                f_der.pack(side="right", padx=12, pady=8)

                ctk.CTkLabel(f_der, text=f"{minutos} minutos", font=("Montserrat", 12, "bold"), text_color=self.C_TEXT_MAIN).pack(anchor="e")
                ctk.CTkLabel(f_der, text=txt_calidad, font=("Arial", 9, "bold"), text_color=color_calidad).pack(anchor="e")

    def _render_vacio(self, mensaje: str):
        f = ctk.CTkFrame(self.scroll_data, fg_color=self.C_PANEL, corner_radius=10, border_width=1, border_color=self.C_BORDER)
        f.pack(fill="x", pady=25, padx=20)
        ctk.CTkLabel(f, text="🔍 " + mensaje, font=("Arial", 11, "italic"), text_color=self.C_TEXT_MUTED).pack(pady=15)

    # =========================================================================
    # EXPORTACIONES PDF & EXCEL INTEGRADAS
    # =========================================================================
    def exportar_pdf_actual(self):
        t = self.tab_selector.get()
        if "SAIME" in t: self.exportar_saime_pdf()
        elif "P2P" in t: self.exportar_p2p_pdf()
        elif "Utilidad" in t: self.exportar_cpp_pdf()
        else: self.exportar_limpieza_pdf()

    def exportar_excel_actual(self):
        t = self.tab_selector.get()
        if "SAIME" in t: self.exportar_saime_excel()
        elif "P2P" in t: self.exportar_p2p_excel()
        elif "Utilidad" in t: self.exportar_cpp_excel()
        else: self.exportar_limpieza_excel()

    # --- MÉTODOS DE EXPORTACIÓN ---
    def exportar_saime_pdf(self):
        if not self.datos_saime:
            messagebox.showwarning("Atención", "No hay datos para exportar.")
            return
        pdf = PDFPlanillaSAIME()
        pdf.add_page()
        pdf.set_font("Helvetica", "", 8)
        nro = 1
        for est in self.datos_saime:
            fecha_f = datetime.strptime(est["fecha_entrada"], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y")
            
            # Fila Titular
            pdf.cell(12, 6, str(nro), 1, 0, "C")
            pdf.cell(22, 6, fecha_f, 1, 0, "C")
            pdf.cell(65, 6, f"{est['nombre']} {est['apellido']}".upper()[:35], 1)
            pdf.cell(12, 6, str(est["edad"]), 1, 0, "C")
            pdf.cell(22, 6, str(est["estado_civil"])[0].upper(), 1, 0, "C")
            pdf.cell(10, 6, str(est["nacionalidad"]).upper(), 1, 0, "C")
            pdf.cell(45, 6, str(est["procedencia"]).upper()[:25], 1)
            pdf.cell(27, 6, str(est["cedula"]), 1, 0, "C")
            pdf.cell(62, 6, str(est["destino"]).upper()[:30], 1)
            pdf.ln()
            nro += 1

            # Fila Acompañante
            if est.get("ac_nombre") and est.get("ac_cedula") and est["ac_nombre"].strip():
                pdf.cell(12, 6, str(nro), 1, 0, "C")
                pdf.cell(22, 6, fecha_f, 1, 0, "C")
                pdf.cell(65, 6, f"{est['ac_nombre']} {est['ac_apellido']}".upper()[:35], 1)
                pdf.cell(12, 6, str(est["ac_edad"]), 1, 0, "C")
                pdf.cell(22, 6, str(est["ac_estado_civil"])[0].upper() if est.get("ac_estado_civil") else "S", 1, 0, "C")
                pdf.cell(10, 6, str(est["ac_nacionalidad"]).upper() if est.get("ac_nacionalidad") else "V", 1, 0, "C")
                pdf.cell(45, 6, str(est["ac_procedencia"]).upper()[:25], 1)
                pdf.cell(27, 6, str(est["ac_cedula"]), 1, 0, "C")
                pdf.cell(62, 6, str(est["ac_destino"]).upper()[:30], 1)
                pdf.ln()
                nro += 1

        ruta = self.obtener_ruta_exportacion(f"planilla_saime_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta)
        os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_saime_excel(self):
        if not self.datos_saime: return
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Planilla SAIME"
        ws.append(["N° PER", "FECHA", "NOMBRE Y APELLIDO", "EDAD", "ESTADO CIVIL", "NAC", "PROCEDENCIA", "CÉDULA", "DESTINO"])
        nro = 1
        for est in self.datos_saime:
            fecha_f = datetime.strptime(est["fecha_entrada"], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y")
            ws.append([nro, fecha_f, f"{est['nombre']} {est['apellido']}".upper(), est["edad"], str(est["estado_civil"])[0].upper(), est["nacionalidad"].upper(), est["procedencia"].upper(), est["cedula"], est["destino"].upper()]); nro += 1
            if est.get("ac_nombre") and est.get("ac_cedula") and est["ac_nombre"].strip():
                ws.append([nro, fecha_f, f"{est['ac_nombre']} {est['ac_apellido']}".upper(), est["ac_edad"], str(est["ac_estado_civil"])[0].upper() if est.get("ac_estado_civil") else "S", est.get("ac_nacionalidad", "V").upper(), est["ac_procedencia"].upper(), est["ac_cedula"], est["ac_destino"].upper()]); nro += 1
        ruta = self.obtener_ruta_exportacion(f"planilla_saime_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_p2p_pdf(self):
        if not self.datos_p2p: return
        pdf = PDFReporteP2P(); pdf.add_page(); pdf.set_font("Helvetica", "", 8)
        total_bs, total_usd = 0.0, 0.0
        for p in self.datos_p2p:
            pdf.cell(30, 6, str(p["fecha"])[:16], 1, 0, "C")
            pdf.cell(28, 6, str(p["hab_codigo"])[:15], 1, 0, "C")
            pdf.cell(38, 6, str(p["concepto"])[:20], 1)
            pdf.cell(50, 6, str(p["cliente_nombre"])[:26], 1)
            pdf.cell(24, 6, str(p["p2p_ci"]), 1, 0, "C")
            pdf.cell(28, 6, str(p["p2p_telefono"]), 1, 0, "C")
            pdf.cell(22, 6, str(p["p2p_referencia"]), 1, 0, "C")
            pdf.cell(32, 6, formatear_bs(p["monto_bs"]), 1, 0, "R")
            pdf.cell(25, 6, f"${p['monto_usd']:.2f}", 1, 0, "R")
            pdf.ln()
            total_bs += p["monto_bs"]
            total_usd += p["monto_usd"]

        pdf.ln(3); pdf.set_font("Helvetica", "B", 9); pdf.set_fill_color(240, 235, 225)
        pdf.cell(220, 7, "TOTAL GENERAL PAGO MÓVIL (P2P):", 1, 0, "L", fill=True)
        pdf.cell(32, 7, formatear_bs(total_bs), 1, 0, "R", fill=True)
        pdf.cell(25, 7, f"${total_usd:.2f}", 1, 1, "R", fill=True)

        ruta = self.obtener_ruta_exportacion(f"reporte_p2p_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_p2p_excel(self):
        if not self.datos_p2p: return
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Auditoría P2P"
        ws.append(["Fecha/Hora", "Habitación", "Concepto", "Cliente/Titular", "C.I. Emisor", "Teléfono", "Referencia", "Monto Bs", "Monto USD", "Turno", "Recepcionista"])
        for p in self.datos_p2p:
            ws.append([p["fecha"], p["hab_codigo"], p["concepto"], p["cliente_nombre"], p["p2p_ci"], p["p2p_telefono"], p["p2p_referencia"], p["monto_bs"], p["monto_usd"], p["turno"], p["recepcionista"]])
        ruta = self.obtener_ruta_exportacion(f"auditoria_p2p_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_cpp_pdf(self):
        if not self.datos_cpp: return
        pdf = PDFReporteUtilidadCPP(); pdf.add_page(); pdf.set_font("Helvetica", "", 7.5)
        tot_ing_usd, tot_cos_usd, tot_util_usd = 0.0, 0.0, 0.0
        tot_ing_bs, tot_util_bs = 0.0, 0.0

        for d in self.datos_cpp:
            pdf.cell(42, 6, str(d["producto"])[:22], 1, 0, "L")
            pdf.cell(14, 6, str(d["unidades_vendidas"]), 1, 0, "C")
            pdf.cell(24, 6, f"${d['costo_prom_usd']:.2f}", 1, 0, "R")
            pdf.cell(24, 6, f"${d['precio_prom_usd']:.2f}", 1, 0, "R")
            pdf.cell(28, 6, f"${d['costo_total_usd']:.2f}", 1, 0, "R")
            pdf.cell(28, 6, f"${d['ingreso_total_usd']:.2f}", 1, 0, "R")
            pdf.cell(38, 6, formatear_bs(d["ingreso_total_bs"]), 1, 0, "R")
            pdf.cell(28, 6, f"${d['utilidad_neta_usd']:.2f}", 1, 0, "R")
            pdf.cell(34, 6, formatear_bs(d["utilidad_neta_bs"]), 1, 0, "R")
            pdf.cell(17, 6, f"{d['margen_pct']:.1f}%", 1, 0, "C")
            pdf.ln()

            tot_ing_usd += d["ingreso_total_usd"]
            tot_cos_usd += d["costo_total_usd"]
            tot_util_usd += d["utilidad_neta_usd"]
            tot_ing_bs += d["ingreso_total_bs"]
            tot_util_bs += d["utilidad_neta_bs"]

        margen_global_pct = (tot_util_usd / tot_ing_usd * 100) if tot_ing_usd > 0 else 0.0

        pdf.ln(3); pdf.set_font("Helvetica", "B", 8); pdf.set_fill_color(240, 235, 225)
        pdf.cell(104, 7, "TOTALES GENERALES EN DÓLARES ($):", 1, 0, "L", fill=True)
        pdf.cell(28, 7, f"${tot_cos_usd:.2f}", 1, 0, "R", fill=True)
        pdf.cell(28, 7, f"${tot_ing_usd:.2f}", 1, 0, "R", fill=True)
        pdf.cell(38, 7, "-", 1, 0, "C", fill=True)
        pdf.cell(28, 7, f"${tot_util_usd:.2f}", 1, 0, "R", fill=True)
        pdf.cell(34, 7, "-", 1, 0, "C", fill=True)
        pdf.cell(17, 7, f"{margen_global_pct:.1f}%", 1, 1, "C", fill=True)

        pdf.set_fill_color(225, 235, 245)
        pdf.cell(104, 7, "TOTALES GENERALES EN BOLÍVARES (BS):", 1, 0, "L", fill=True)
        pdf.cell(28, 7, "-", 1, 0, "C", fill=True)
        pdf.cell(28, 7, "-", 1, 0, "C", fill=True)
        pdf.cell(38, 7, formatear_bs(tot_ing_bs), 1, 0, "R", fill=True)
        pdf.cell(28, 7, "-", 1, 0, "C", fill=True)
        pdf.cell(34, 7, formatear_bs(tot_util_bs), 1, 0, "R", fill=True)
        pdf.cell(17, 7, f"{margen_global_pct:.1f}%", 1, 1, "C", fill=True)

        ruta = self.obtener_ruta_exportacion(f"utilidad_minibar_cpp_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_cpp_excel(self):
        if not self.datos_cpp: return
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Utilidad Mini-Bar CPP"
        ws.append(["Producto", "Cant", "Costo CPP ($)", "Precio Venta ($)", "Total Costo ($)", "Total Venta ($)", "Total Venta (Bs)", "Utilidad Neta ($)", "Utilidad Neta (Bs)", "Margen %"])
        for d in self.datos_cpp:
            ws.append([d["producto"], d["unidades_vendidas"], d["costo_prom_usd"], d["precio_prom_usd"], d["costo_total_usd"], d["ingreso_total_usd"], d["ingreso_total_bs"], d["utilidad_neta_usd"], d["utilidad_neta_bs"], f"{d['margen_pct']:.1f}%"])
        ruta = self.obtener_ruta_exportacion(f"utilidad_minibar_cpp_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_limpieza_pdf(self):
        if not self.datos_limpieza:
            messagebox.showwarning("Atención", "No hay registros de limpieza para exportar.")
            return

        pdf = PDFReporteLimpieza(); pdf.add_page(); pdf.set_font("Helvetica", "", 8)
        for l in self.datos_limpieza:
            pdf.cell(32, 6, str(l["fecha_limpia"])[:16], 1, 0, "C")
            pdf.cell(25, 6, str(l["hab_codigo"]), 1, 0, "C")
            pdf.cell(60, 6, str(l["camarera"])[:30], 1, 0, "L")
            pdf.cell(25, 6, str(l["turno"]), 1, 0, "C")
            pdf.cell(50, 6, str(l["fecha_sucia"])[:16] if l["fecha_sucia"] else "N/A", 1, 0, "C")
            pdf.cell(45, 6, f"{l['duracion_minutos']} minutos", 1, 0, "C")
            
            estado_cal = "ÓPTIMO" if l["duracion_minutos"] <= 45 else "DEMORADO"
            pdf.cell(40, 6, estado_cal, 1, 0, "C")
            pdf.ln()

        ruta = self.obtener_ruta_exportacion(f"control_camareras_sla_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')

    def exportar_limpieza_excel(self):
        if not self.datos_limpieza: return
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Control Camareras SLA"
        ws.append(["Fecha Limpia", "Habitación", "Camarera", "Turno", "Hora Sucia", "Duración Minutos", "Evaluación SLA"])

        for l in self.datos_limpieza:
            eval_txt = "ÓPTIMO" if l["duracion_minutos"] <= 45 else "DEMORADO"
            ws.append([l["fecha_limpia"], l["hab_codigo"], l["camarera"], l["turno"], l["fecha_sucia"] or "N/A", l["duracion_minutos"], eval_txt])

        ruta = self.obtener_ruta_exportacion(f"control_camareras_sla_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(ruta); os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')