"""
===============================================================================
SISTEMA DE GESTIÓN DE MOTELES - MOTEL PRO v2.0
Módulo: modules/mantenimiento.py (CONTROL DE AVERÍAS, GASTOS Y DAÑOS A HUÉSPEDES)
===============================================================================
"""

import os
import openpyxl
import customtkinter as ctk
from tkinter import messagebox
from datetime import datetime
from fpdf import FPDF

# IMPORTACIONES COMPLETAS Y VERIFICADAS
from database.db_manager import (
    obtener_habitaciones, enviar_habitacion_mantenimiento,
    finalizar_mantenimiento_habitacion, obtener_ordenes_mantenimiento,
    registrar_cobro_dano, registrar_auditoria, formatear_bs, formatear_usd,
    obtener_tasa_bcv, obtener_ruta_recurso
)


def centrar_ventana(win, ancho: int, alto: int):
    """Calcula y posiciona la ventana en el centro del monitor."""
    win.update_idletasks()
    pantalla_ancho = win.winfo_screenwidth()
    pantalla_alto = win.winfo_screenheight()
    x = max(0, (pantalla_ancho - ancho) // 2)
    y = max(0, (pantalla_alto - alto) // 2)
    win.geometry(f"{ancho}x{alto}+{x}+{y}")


# =============================================================================
# PDF: HISTORIAL DE MANTENIMIENTO Y GASTOS ($ Y BS)
# =============================================================================
class PDFReporteMantenimiento(FPDF):
    def __init__(self):
        super().__init__(orientation='L', unit='mm', format='A4')  # 277 mm de área útil
        self.set_margins(10, 10, 10)
        self.set_auto_page_break(auto=True, margin=12)

    def header(self):
        logo = obtener_ruta_recurso("logo_hotel.png")
        if os.path.exists(logo):
            self.image(logo, x=10, y=6, w=20)

        self.set_font("Arial", "B", 12)
        self.cell(0, 5, "INVERSIONES SAIBABA C.A. - MOTEL EL EDEN", 0, 1, "C")
        self.set_font("Arial", "B", 9)
        self.cell(0, 4.5, "RIF: J-30250227-6 | HISTORIAL DE MANTENIMIENTO, AVERÍAS Y COSTOS DE REPARACIÓN", 0, 1, "C")
        self.set_y(26)

        self.set_fill_color(220, 230, 242)
        self.set_font("Arial", "B", 8)
        self.cell(28, 7, "FECHA REPORTE", 1, 0, "C", fill=True)
        self.cell(22, 7, "HABITACIÓN", 1, 0, "C", fill=True)
        self.cell(35, 7, "TIPO DE AVERÍA", 1, 0, "C", fill=True)
        self.cell(72, 7, "DESCRIPCIÓN DE LA FALLA", 1, 0, "L", fill=True)
        self.cell(38, 7, "TÉCNICO ASIGNADO", 1, 0, "L", fill=True)
        self.cell(24, 7, "COSTO ($)", 1, 0, "R", fill=True)
        self.cell(36, 7, "COSTO (BS)", 1, 0, "R", fill=True)
        self.cell(22, 7, "ESTADO", 1, 0, "C", fill=True)
        self.ln()

    def footer(self):
        self.set_y(-12)
        self.set_font("Arial", "I", 8)
        self.cell(0, 8, f"Auditoría de Mantenimiento e Infraestructura - Página {self.page_no()}", 0, 0, "C")


class FrameMantenimiento(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, fg_color="#191715")
        self.controller = controller

        # Encabezado Superior
        self.frame_top = ctk.CTkFrame(self, fg_color="#23201C", corner_radius=12, border_width=1, border_color="#36322C")
        self.frame_top.pack(fill="x", padx=15, pady=10)

        ctk.CTkLabel(self.frame_top, text="🛠️ Control de Mantenimiento, Averías y Daños", font=("Georgia", 16, "bold"), text_color="#D4A343").pack(side="left", padx=15, pady=10)

        ctk.CTkButton(
            self.frame_top, text="📋 Reporte Gastos PDF", fg_color="#3E342B", hover_color="#524539",
            text_color="#F4EFE6", font=("Arial", 11, "bold"), command=self.exportar_gastos_pdf
        ).pack(side="right", padx=6)

        ctk.CTkButton(
            self.frame_top, text="⚡ Cobrar Daño a Huésped", fg_color="#78281F", hover_color="#943126",
            text_color="#F4EFE6", font=("Arial", 11, "bold"), command=self.modal_cobro_dano
        ).pack(side="right", padx=6)

        ctk.CTkButton(
            self.frame_top, text="➕ Poner en Mantenimiento", fg_color="#D4A343", hover_color="#B8892E",
            text_color="#191715", font=("Arial", 11, "bold"), command=self.modal_nueva_orden
        ).pack(side="right", padx=6)

        # Contenedor de Órdenes
        self.scroll_ordenes = ctk.CTkScrollableFrame(self, fg_color="#23201C", label_text="Habitaciones Fuera de Servicio y Reparaciones Activas")
        self.scroll_ordenes.pack(fill="both", expand=True, padx=15, pady=10)

        self.cargar_ordenes()

    def cargar_ordenes(self):
        for w in self.scroll_ordenes.winfo_children():
            w.destroy()

        ordenes = obtener_ordenes_mantenimiento(solo_activas=True)
        if not ordenes:
            ctk.CTkLabel(self.scroll_ordenes, text="Todas las habitaciones están operativas. No hay averías activas.", font=("Arial", 12, "italic"), text_color="#A89F91").pack(pady=40)
            return

        for ord_item in ordenes:
            card = ctk.CTkFrame(self.scroll_ordenes, fg_color="#2D2924", corner_radius=10, border_width=1, border_color="#4A4F55")
            card.pack(fill="x", pady=5, padx=5)

            f_txt = f"HABITACIÓN: {ord_item['hab_codigo']} | Avería: {ord_item['tipo_averia']}\nDescripción: {ord_item['descripcion']}\nTécnico: {ord_item['tecnico_responsable']} | Costo Est: ${ord_item['costo_reparacion_usd']:.2f} ({formatear_bs(ord_item['costo_reparacion_bs'])})\nFecha Reporte: {ord_item['fecha_inicio']}"
            ctk.CTkLabel(card, text=f_txt, font=("Arial", 11), text_color="#F4EFE6", justify="left").pack(side="left", padx=12, pady=10)

            btn_fin = ctk.CTkButton(
                card, text="✅ Reparada (Habilitar)", fg_color="#1E5F38", hover_color="#2E7D32", width=140,
                command=lambda oid=ord_item["id"], hcod=ord_item["hab_codigo"]: self.liberar_habitacion(oid, hcod)
            )
            btn_fin.pack(side="right", padx=15, pady=10)

    def liberar_habitacion(self, orden_id: int, hab_codigo: str):
        usr_act = self.controller.usuario_actual["usuario"]
        rol_act = self.controller.usuario_actual["rol"]

        if messagebox.askyesno("Confirmación", f"¿Confirmar que la habitación {hab_codigo} fue reparada? Pasará a estado 'Limpia' (Verde)."):
            if finalizar_mantenimiento_habitacion(orden_id, hab_codigo):
                registrar_auditoria(usr_act, rol_act, "FIN_MANTENIMIENTO", f"Mantenimiento #{orden_id} finalizado. Habitación {hab_codigo} habilitada a 'Limpia'.")
                messagebox.showinfo("Éxito", f"Habitación {hab_codigo} habilitada con éxito.")
                self.cargar_ordenes()

    def modal_nueva_orden(self):
        win = ctk.CTkToplevel(self)
        win.title("Poner Habitación en Mantenimiento")
        centrar_ventana(win, 450, 440)
        win.grab_set()

        habs_disponibles = [h["codigo"] for h in obtener_habitaciones()]

        ctk.CTkLabel(win, text="Orden de Reparación / Avería", font=("Georgia", 14, "bold"), text_color="#D4A343").pack(pady=10)

        ctk.CTkLabel(win, text="Seleccionar Habitación:").pack(anchor="w", padx=20)
        c_hab = ctk.CTkComboBox(win, values=habs_disponibles)
        c_hab.pack(fill="x", padx=20, pady=4)

        ctk.CTkLabel(win, text="Tipo de Avería:").pack(anchor="w", padx=20)
        c_tipo = ctk.CTkComboBox(win, values=["Plomería / Baño", "Aire Acondicionado", "Jacuzzi / Bomba", "Electricidad / Iluminación", "Cerrajería / Puerta", "Televisor / Cable", "Carpintería / Cama"])
        c_tipo.pack(fill="x", padx=20, pady=4)

        e_desc = ctk.CTkEntry(win, placeholder_text="Descripción de la falla")
        e_desc.pack(fill="x", padx=20, pady=6)

        e_tec = ctk.CTkEntry(win, placeholder_text="Técnico / Encargado de Reparación")
        e_tec.pack(fill="x", padx=20, pady=6)

        e_cost = ctk.CTkEntry(win, placeholder_text="Costo Estimado de Reparación ($)")
        e_cost.pack(fill="x", padx=20, pady=6)

        def guardar():
            h, t, d, tec, c = c_hab.get(), c_tipo.get(), e_desc.get().strip(), e_tec.get().strip(), e_cost.get().strip()
            if not all([h, t, d, tec, c]):
                messagebox.showwarning("Atención", "Complete todos los campos.")
                return
            try:
                c_flt = float(c.replace(",", "."))
                usr = self.controller.usuario_actual["usuario"]
                rol = self.controller.usuario_actual["rol"]
                if enviar_habitacion_mantenimiento(h, t, d, tec, c_flt, usr):
                    registrar_auditoria(usr, rol, "INICIO_MANTENIMIENTO", f"Habitación {h} puesta en Mantenimiento por '{t}'. Técnico: {tec}. Costo est: ${c_flt:.2f}.")
                    messagebox.showinfo("Éxito", f"Habitación {h} fuera de servicio por mantenimiento.")
                    win.destroy()
                    self.cargar_ordenes()
            except ValueError:
                messagebox.showerror("Error", "El costo debe ser numérico.")

        ctk.CTkButton(win, text="Confirmar y Bloquear Habitación", fg_color="#78281F", hover_color="#943126", command=guardar).pack(fill="x", padx=20, pady=15)

    # =========================================================================
    # MODAL: REPORTE Y COBRO DE DAÑOS A HUÉSPED (CON NOTIFICACIÓN Y AUDITORÍA)
    # =========================================================================
    def modal_cobro_dano(self):
        """Modal profesional para cobrar daños con C.I., cálculo BCV y soporte Pago Móvil."""
        win = ctk.CTkToplevel(self)
        win.title("Reporte y Cobro de Daños a Huésped")
        centrar_ventana(win, 480, 560)
        win.grab_set()

        tasa = obtener_tasa_bcv()

        ctk.CTkLabel(win, text="Cobro de Daños Ocasionados por Huésped", font=("Georgia", 15, "bold"), text_color="#D4A343").pack(pady=(15, 2))
        ctk.CTkLabel(win, text=f"Tasa BCV del Día: {tasa:.2f} Bs / USD", font=("Arial", 10, "italic"), text_color="#A89F91").pack(pady=(0, 8))

        f_dano = ctk.CTkFrame(win, fg_color="#23201C", corner_radius=10)
        f_dano.pack(fill="both", expand=True, padx=15, pady=5)

        ctk.CTkLabel(f_dano, text="Habitación:", font=("Arial", 11, "bold"), text_color="#A89F91").pack(anchor="w", padx=15, pady=(8, 1))
        e_hab = ctk.CTkEntry(f_dano, placeholder_text="Ej: VIP-01", fg_color="#2D2924")
        e_hab.pack(fill="x", padx=15, pady=2)

        ctk.CTkLabel(f_dano, text="Nombre del Huésped Responsable:", font=("Arial", 11, "bold"), text_color="#A89F91").pack(anchor="w", padx=15, pady=(5, 1))
        e_cli = ctk.CTkEntry(f_dano, placeholder_text="Nombre y Apellido", fg_color="#2D2924")
        e_cli.pack(fill="x", padx=15, pady=2)

        ctk.CTkLabel(f_dano, text="Cédula de Identidad (C.I.):", font=("Arial", 11, "bold"), text_color="#A89F91").pack(anchor="w", padx=15, pady=(5, 1))
        e_ci_cli = ctk.CTkEntry(f_dano, placeholder_text="Sólo números (Ej: 20000000)", fg_color="#2D2924")
        e_ci_cli.pack(fill="x", padx=15, pady=2)

        ctk.CTkLabel(f_dano, text="Concepto del Daño / Avería:", font=("Arial", 11, "bold"), text_color="#A89F91").pack(anchor="w", padx=15, pady=(5, 1))
        c_dano = ctk.CTkComboBox(
            f_dano,
            values=[
                "Sábanas / Lencería Quemada o Rota", "Control de TV Dañado / Extraviado",
                "Toallas Dañadas / Extraviadas", "Vidrio / Espejo Roto",
                "Daño a Jacuzzi / Grifería", "Control de Aire Dañado", "Manchas / Daños en Colchón"
            ]
        )
        c_dano.pack(fill="x", padx=15, pady=2)

        ctk.CTkLabel(f_dano, text="Monto a Cobrar ($ USD):", font=("Arial", 11, "bold"), text_color="#D4A343").pack(anchor="w", padx=15, pady=(5, 1))
        e_monto = ctk.CTkEntry(f_dano, placeholder_text="Ej: 25.00", fg_color="#2D2924")
        e_monto.pack(fill="x", padx=15, pady=2)

        lbl_bs_calc = ctk.CTkLabel(f_dano, text="Equivalente: 0,00 Bs", font=("Arial", 11, "bold"), text_color="#7DCEA0")
        lbl_bs_calc.pack(anchor="w", padx=15, pady=2)

        def actualizar_monto_bs(event=None):
            try:
                m_val = float(e_monto.get().replace(",", "."))
                lbl_bs_calc.configure(text=f"Equivalente: {formatear_bs(m_val * tasa)}")
            except ValueError:
                lbl_bs_calc.configure(text="Equivalente: 0,00 Bs")

        e_monto.bind("<KeyRelease>", actualizar_monto_bs)

        ctk.CTkLabel(f_dano, text="Forma de Pago:", font=("Arial", 11, "bold"), text_color="#A89F91").pack(anchor="w", padx=15, pady=(5, 1))
        c_met = ctk.CTkComboBox(
            f_dano, values=["EFECTIVO USD", "EFECTIVO BS", "PAGO MOVIL", "PUNTO DE VENTA"],
            command=lambda v: toggle_p2p_dano(v)
        )
        c_met.pack(fill="x", padx=15, pady=2)

        frame_p2p_d = ctk.CTkFrame(f_dano, fg_color="transparent")
        e_ref_d = ctk.CTkEntry(frame_p2p_d, placeholder_text="Ref (4 dígitos)", fg_color="#2D2924")
        e_ref_d.pack(side="left", padx=2, fill="x", expand=True)
        e_ci_p2p = ctk.CTkEntry(frame_p2p_d, placeholder_text="C.I. Emisor", fg_color="#2D2924")
        e_ci_p2p.pack(side="left", padx=2, fill="x", expand=True)
        e_tel_d = ctk.CTkEntry(frame_p2p_d, placeholder_text="Teléfono Emisor", fg_color="#2D2924")
        e_tel_d.pack(side="left", padx=2, fill="x", expand=True)

        def toggle_p2p_dano(metodo):
            if metodo == "PAGO MOVIL":
                frame_p2p_d.pack(fill="x", padx=15, pady=5)
            else:
                frame_p2p_d.pack_forget()

        def cobrar():
            h = e_hab.get().strip()
            cli = e_cli.get().strip()
            ci = e_ci_cli.get().strip()
            con = c_dano.get()
            m = e_monto.get().strip()
            met = c_met.get()

            if not all([h, cli, ci, con, m]):
                messagebox.showwarning("Atención", "Complete todos los campos del huésped y el daño.")
                return

            if not ci.isdigit():
                messagebox.showwarning("Atención", "La Cédula de Identidad debe ser numérica.")
                return

            ref, ci_emisor, tel = None, None, None
            if met == "PAGO MOVIL":
                ref = e_ref_d.get().strip()
                ci_emisor = e_ci_p2p.get().strip()
                tel = e_tel_d.get().strip()
                if not (ref and ci_emisor and tel):
                    messagebox.showerror("Error", "Complete los datos del Pago Móvil.")
                    return

            try:
                m_flt = float(m.replace(",", "."))
                usr = self.controller.usuario_actual["usuario"]
                rol = self.controller.usuario_actual["rol"]

                if registrar_cobro_dano(h, cli, ci, con, m_flt, met, usr, ref, ci_emisor, tel):
                    p2p_txt = f" [Ref P2P: {ref}]" if met == "PAGO MOVIL" else ""
                    
                    # 1. Auditoría interna
                    registrar_auditoria(
                        usr, rol, "COBRO_DANO",
                        f"Cobro de daño en Hab {h} a {cli} (C.I: {ci}): '{con}' por ${m_flt:.2f} ({formatear_bs(m_flt * tasa)}) [{met}{p2p_txt}]."
                    )

                    # 2. Notificación Multicanal en segundo plano
                    try:
                        from modules.notificaciones import disparar_notificacion_multicanal
                        asunto_d = f"⚡ DAÑO COBRADO EN HAB {h}"
                        msg_d = (
                            f"🏨 <b>ALERTA DE DAÑO EN HABITACIÓN</b>\n"
                            f"🚪 <b>Habitación:</b> {h}\n"
                            f"👤 <b>Huésped:</b> {cli} (C.I: {ci})\n"
                            f"⚠️ <b>Concepto:</b> {con}\n"
                            f"💰 <b>Monto Cobrado:</b> ${m_flt:.2f} ({formatear_bs(m_flt * tasa)}) [{met}{p2p_txt}]\n"
                            f"👮 <b>Recepcionista:</b> {usr}\n"
                        )
                        disparar_notificacion_multicanal(asunto_d, msg_d)
                    except Exception as e:
                        print(f"Error notificando daño: {e}")

                    messagebox.showinfo("Éxito", "Cobro de daño registrado e ingresado a la caja del turno.")
                    win.destroy()
            except ValueError:
                messagebox.showerror("Error", "Monto inválido.")

        ctk.CTkButton(
            win, text="💳 Cobrar y Registrar Daño a Caja", fg_color="#D4A343", hover_color="#B8892E",
            text_color="#191715", font=("Arial", 12, "bold"), command=cobrar
        ).pack(fill="x", padx=15, pady=12)

    def exportar_gastos_pdf(self):
        ordenes = obtener_ordenes_mantenimiento(solo_activas=False)
        if not ordenes:
            messagebox.showwarning("Atención", "No hay registros de reparaciones para exportar.")
            return

        pdf = PDFReporteMantenimiento()
        pdf.add_page()
        pdf.set_font("Arial", "", 8)

        tot_usd, tot_bs = 0.0, 0.0
        for o in ordenes:
            pdf.cell(28, 6, str(o["fecha_inicio"])[:16], 1, 0, "C")
            pdf.cell(22, 6, str(o["hab_codigo"]), 1, 0, "C")
            pdf.cell(35, 6, str(o["tipo_averia"])[:18], 1, 0, "L")
            pdf.cell(72, 6, str(o["descripcion"])[:42], 1, 0, "L")
            pdf.cell(38, 6, str(o["tecnico_responsable"])[:20], 1, 0, "L")
            pdf.cell(24, 6, f"${o['costo_reparacion_usd']:.2f}", 1, 0, "R")
            pdf.cell(36, 6, formatear_bs(o["costo_reparacion_bs"]), 1, 0, "R")
            pdf.cell(22, 6, str(o["estado"]), 1, 0, "C")
            pdf.ln()

            tot_usd += o["costo_reparacion_usd"]
            tot_bs += o["costo_reparacion_bs"]

        # Resumen Final Bimonetario
        pdf.ln(3)
        pdf.set_font("Arial", "B", 8.5)
        pdf.set_fill_color(220, 230, 242)
        pdf.cell(195, 7, "TOTAL GASTOS DE REPARACIÓN E INFRAESTRUCTURA:", 1, 0, "L", fill=True)
        pdf.cell(24, 7, f"${tot_usd:.2f}", 1, 0, "R", fill=True)
        pdf.cell(36, 7, formatear_bs(tot_bs), 1, 0, "R", fill=True)
        pdf.cell(22, 7, "", 1, 1, "C", fill=True)

        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        export_dir = os.path.join(base_dir, "reportes_exportados")
        os.makedirs(export_dir, exist_ok=True)

        ruta = os.path.join(export_dir, f"reporte_mantenimiento_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        pdf.output(ruta)
        os.system(f'start "" "{ruta}"' if os.name == 'nt' else f'open "{ruta}"')